"""UX-07: сопоставление ответов реестра с журналом, импорт файла-ответа, диалог запроса."""
import os
import sys
import tempfile
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))
os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')

import pytest

from db.database import DatabaseManager
from db.schema import create_schema
from db.exam_journal_repo import JournalRecord
from journal.journal_manager import JournalManager
from journal.registry_sync import (
    merge_registry_records, api_records_from_rows, is_api_xlsx_headers,
    normalize_date, setid_from_filename, format_snils, describe_merge,
)
from utils.crypto import clear_caches

# Synthetic SNILS with a valid checksum (not a real person)
SNILS_A = "112-233-445 95"
SNILS_B = "123-456-789 64"


def _jr(snils=SNILS_A, program="1", protocol="П-1", date="01.02.2026",
        set_id="", base_no="", status="pending", uid=None):
    return JournalRecord(
        uuid=uid or f"u-{snils}-{program}-{protocol}", send_date="", set_id=set_id, xml_file="",
        last_name="Тестов", first_name="Тест", middle_name="Тестович", snils=snils,
        position="Инженер", program_id=program, program_title="", exam_date=date,
        protocol=protocol, result="Удовлетворительно", base_no=base_no, status=status)


def _api(snils=SNILS_A, program="1", protocol="П-1", date="2026-02-01T00:00:00", base_no="R-1",
         passed="true"):
    return {"Snils": snils, "learnProgramId": program, "ProtocolNumber": protocol, "Date": date,
            "baseNo": base_no, "LastName": "Тестов", "FirstName": "Тест", "MiddleName": "Тестович",
            "Position": "Инженер", "LearnProgramTitle": "", "isPassed": passed}


# ── pure logic ──────────────────────────────────────────────

class TestHelpers:
    @pytest.mark.parametrize("raw,expected", [
        ("2025-09-26T00:00:00", "26.09.2025"),
        ("2025-09-26", "26.09.2025"),
        ("26.09.2025", "26.09.2025"),
        ("26.09.2025 10:00", "26.09.2025"),
        ("", ""),
        (None, ""),
    ])
    def test_normalize_date(self, raw, expected):
        assert normalize_date(raw) == expected

    def test_setid_from_filename(self):
        assert setid_from_filename(r"C:\x\123456.xlsx") == "123456"
        assert setid_from_filename("отчёт.xlsx") == ""

    def test_format_snils(self):
        assert format_snils("11223344595") == SNILS_A


class TestMerge:
    def test_match_by_protocol_updates_base_no_and_fills_set_id(self):
        res = merge_registry_records([_api(date="")], [_jr()], set_id="777")
        assert res.updates == [("u-%s-1-П-1" % SNILS_A, "R-1", "777")]
        assert res.new_records == [] and res.already == 0

    def test_match_by_date_when_protocol_differs(self):
        res = merge_registry_records([_api(protocol="другой")], [_jr()])
        assert res.updated == 1

    def test_snils_format_and_program_float_tolerated(self):
        rec = _api(snils="11223344595", program="1.0")
        assert merge_registry_records([rec], [_jr()]).updated == 1

    def test_already_received_same_number_is_not_rewritten(self):
        existing = _jr(base_no="R-1", status="received", set_id="777")
        res = merge_registry_records([_api()], [existing], set_id="777")
        assert res.updated == 0 and res.already == 1

    def test_existing_set_id_is_not_overwritten(self):
        res = merge_registry_records([_api()], [_jr(set_id="111")], set_id="222")
        assert res.updates[0][2] == ""

    def test_other_program_becomes_new_record(self):
        res = merge_registry_records([_api(program="2")], [_jr()], set_id="777")
        assert res.updated == 0
        assert len(res.new_records) == 1
        new = res.new_records[0]
        assert (new.snils, new.program_id, new.exam_date, new.set_id, new.status) == \
            (SNILS_A, "2", "01.02.2026", "777", "received")

    def test_failed_result_and_duplicates_in_response(self):
        res = merge_registry_records([_api(program="9", passed="false")] * 2, [])
        assert len(res.new_records) == 1
        assert res.new_records[0].result == "Неудовлетворительно"

    def test_invalid_rows_are_skipped(self):
        res = merge_registry_records([_api(snils="123"), _api(program="")], [_jr()])
        assert res.skipped == 2 and res.updated == 0

    def test_describe(self):
        res = merge_registry_records([_api(), _api(program="2")], [_jr()])
        assert describe_merge(res) == "Обновлено записей журнала: 1; нет в журнале: 1"
        assert "добавлено новых: 1" in describe_merge(res, added=1)


class TestXlsxResponseFormat:
    def test_roundtrip_with_transfer_tab_export(self, tmp_path):
        """Файл, сохранённый «Передача данных → Запрос по SetId», распознаётся и читается."""
        from openpyxl import load_workbook
        from api.mintrud_api import export_records_to_xlsx
        fp = tmp_path / "555.xlsx"
        ok, _ = export_records_to_xlsx([_api(), _api(snils=SNILS_B, program="18", base_no="R-2")], str(fp))
        assert ok
        ws = load_workbook(fp).active
        headers = [c.value for c in ws[1]]
        assert is_api_xlsx_headers(headers)
        recs = api_records_from_rows(headers, ws.iter_rows(min_row=2, values_only=True))
        assert [r["baseNo"] for r in recs] == ["R-1", "R-2"]
        assert recs[0]["Date"] == "01.02.2026"
        assert setid_from_filename(str(fp)) == "555"

    def test_journal_headers_are_not_api_format(self):
        from utils.table_models import JOURNAL_FIELD_NAMES
        assert not is_api_xlsx_headers(JOURNAL_FIELD_NAMES)


# ── DB-backed ───────────────────────────────────────────────

@pytest.fixture
def db():
    clear_caches()
    with tempfile.NamedTemporaryFile(suffix='.db', delete=False) as f:
        db_path = f.name
    db = DatabaseManager.get_instance(db_path)
    db.initialize()
    create_schema()
    yield db
    db.close_all()
    if os.path.exists(db_path):
        os.remove(db_path)


class TestJournalManagerRegistry:
    def test_apply_updates_persist(self, db, tmp_path):
        jm = JournalManager(str(tmp_path))
        jm.add_journal_records_directly([_jr(uid="a"), _jr(snils=SNILS_B, uid="b", set_id="111")])
        res = jm.match_registry_records([_api(), _api(snils=SNILS_B, base_no="R-2")], set_id="777")
        assert jm.apply_registry_updates(res) == 2
        by_uuid = {r.uuid: r for r in jm.get_all_records()}
        assert (by_uuid["a"].base_no, by_uuid["a"].status, by_uuid["a"].set_id) == ("R-1", "received", "777")
        assert (by_uuid["b"].base_no, by_uuid["b"].set_id) == ("R-2", "111")

    def test_new_record_gets_program_title_from_reference(self, db, tmp_path):
        jm = JournalManager(str(tmp_path))
        res = jm.match_registry_records([_api(program="9")])
        assert res.new_records[0].program_title == JournalManager._get_program_title("9")


# ── UI ──────────────────────────────────────────────────────

@pytest.fixture(scope="module")
def qapp():
    from PySide6.QtWidgets import QApplication
    return QApplication.instance() or QApplication([])


class TestQueryDialog:
    def test_validate_values(self):
        from tabs.registry_query_dialog import validate_query_value, MODE_SETID, MODE_SNILS
        assert validate_query_value(MODE_SETID, " 12345 ") == ("12345", None)
        assert validate_query_value(MODE_SETID, "SET-1")[1]
        assert validate_query_value(MODE_SETID, "")[1] == "Введите SetId"
        assert validate_query_value(MODE_SNILS, "11223344595") == (SNILS_A, None)
        assert validate_query_value(MODE_SNILS, "112-233-445 96")[1] == "Неверная контрольная сумма СНИЛС"
        assert validate_query_value(MODE_SNILS, "123")[1] == "СНИЛС должен содержать 11 цифр"

    def test_default_is_setid_and_accept(self, qapp):
        from tabs.registry_query_dialog import RegistryQueryDialog, MODE_SETID, MODE_SNILS
        d = RegistryQueryDialog()
        assert d.mode() == MODE_SETID
        d.value_input.setText("98765")
        d._on_query()
        assert d.result() == RegistryQueryDialog.DialogCode.Accepted and d.value() == "98765"

        d2 = RegistryQueryDialog()
        d2.snils_radio.setChecked(True)
        assert d2.mode() == MODE_SNILS and d2.value_input.maxLength() == 14
        d2.value_input.setText("123")
        d2._on_query()
        assert d2.result() != RegistryQueryDialog.DialogCode.Accepted
        assert d2.error_label.text()

    def test_worker_calls_api_by_mode(self, qapp, monkeypatch):
        import api.mintrud_api as api
        from tabs.registry_query_dialog import RegistryQueryWorker, MODE_SETID, MODE_SNILS
        calls = []
        monkeypatch.setattr(api, "get_by_set_id", lambda k, v, **kw: calls.append(("setid", v)) or
                            {"success": True, "records": [_api()]})
        monkeypatch.setattr(api, "get_by_snils", lambda k, v, **kw: calls.append(("snils", v)) or
                            {"success": True, "records": []})
        got = []
        for mode, value in ((MODE_SETID, "1"), (MODE_SNILS, SNILS_A)):
            w = RegistryQueryWorker(mode, value, "k" * 32, None)
            w.finished.connect(got.append)
            w.run()
        assert calls == [("setid", "1"), ("snils", SNILS_A)]
        assert got[0]["mode"] == MODE_SETID and got[0]["value"] == "1"


def test_records_phrase_plural():
    from tabs.exam_journal_tab import _records_phrase
    assert [_records_phrase(n) for n in (1, 2, 5, 11, 21, 22, 112)] == [
        "1 запись", "2 записи", "5 записей", "11 записей", "21 запись", "22 записи", "112 записей"]


class TestJournalTabMerge:
    def _tab(self, tmp_path, monkeypatch, answer):
        from PySide6.QtWidgets import QMessageBox
        import tabs.exam_journal_tab as jt
        monkeypatch.setattr(QMessageBox, "question", lambda *a, **k: answer)
        toasts = []
        monkeypatch.setattr(jt.Toast, "success", lambda parent, msg, *a, **k: toasts.append(msg))
        monkeypatch.setattr(jt.Toast, "info", lambda parent, msg, *a, **k: toasts.append(msg))
        jm = JournalManager(str(tmp_path))
        jm.add_journal_records_directly([_jr(uid="a")])
        return jt.ExamJournalTab(jm, str(tmp_path)), jm, toasts

    def test_file_import_updates_and_adds(self, qapp, db, tmp_path, monkeypatch):
        from PySide6.QtWidgets import QMessageBox, QFileDialog
        from api.mintrud_api import export_records_to_xlsx
        tab, jm, toasts = self._tab(tmp_path, monkeypatch, QMessageBox.StandardButton.Yes)
        fp = tmp_path / "4242.xlsx"
        export_records_to_xlsx([_api(), _api(snils=SNILS_B, program="18", base_no="R-2")], str(fp))
        monkeypatch.setattr(QFileDialog, "getOpenFileName", lambda *a, **k: (str(fp), ""))
        tab._import_from_excel()

        recs = {r.snils: r for r in jm.get_all_records()}
        assert len(recs) == 2, "существующая запись обновлена, а не продублирована"
        assert (recs[SNILS_A].base_no, recs[SNILS_A].status, recs[SNILS_A].set_id) == ("R-1", "received", "4242")
        assert recs[SNILS_B].program_id == "18" and recs[SNILS_B].set_id == "4242"
        assert toasts[-1] == "Обновлено записей журнала: 1; добавлено новых: 1"

    def test_query_result_declined_new_records(self, qapp, db, tmp_path, monkeypatch):
        from PySide6.QtWidgets import QMessageBox
        tab, jm, toasts = self._tab(tmp_path, monkeypatch, QMessageBox.StandardButton.No)
        tab._on_registry_query_result({"success": True, "mode": "snils", "value": SNILS_A,
                                       "records": [_api(), _api(program="2")]})
        recs = jm.get_all_records()
        assert len(recs) == 1 and recs[0].base_no == "R-1"
        assert recs[0].set_id == "", "запрос по СНИЛС не проставляет SetId"
        assert toasts[-1] == "Обновлено записей журнала: 1; нет в журнале: 1"

    def test_query_error_shows_message(self, qapp, db, tmp_path, monkeypatch):
        from PySide6.QtWidgets import QMessageBox
        import tabs.exam_journal_tab as jt
        tab, jm, toasts = self._tab(tmp_path, monkeypatch, QMessageBox.StandardButton.No)
        shown = []
        monkeypatch.setattr(jt, "safe_message_box", lambda *a, **k: shown.append(a[3]))
        tab._on_registry_query_result({"success": False, "error": "нет связи", "mode": "setid"})
        assert shown == ["нет связи"]
